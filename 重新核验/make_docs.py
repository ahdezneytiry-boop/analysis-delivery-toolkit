"""使用原生Word公式生成两份模型文档；读取真实验证和回测结果。"""
import solve as m
import json,re,copy
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def atom(txt):
    r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=txt;r.append(t);return r

def mathparts(text):
    # 变量下标、上标转原生OMML，正文不残留LaTeX源码。
    out=[];pos=0
    pat=re.compile(r'([A-Za-zΑ-ωΣλδηπ]+)(?:_\{([^{}]+)\})?(?:\^\{([^{}]+)\})?')
    for mm in pat.finditer(text):
        if mm.start()>pos:out.append(atom(text[pos:mm.start()]))
        name,sub,sup=mm.groups()
        if sub or sup:
            obj=OxmlElement('m:sSubSup' if sub and sup else ('m:sSub' if sub else 'm:sSup'));e=OxmlElement('m:e');e.append(atom(name));obj.append(e)
            for tag,value in [('sub',sub),('sup',sup)]:
                if value:node=OxmlElement('m:'+tag);node.append(atom(value));obj.append(node)
            out.append(obj)
        else:out.append(atom(name))
        pos=mm.end()
    if pos<len(text):out.append(atom(text[pos:]))
    return out

def eq(doc,text):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6)
    om=OxmlElement('m:oMathPara');oo=OxmlElement('m:oMath')
    for part in mathparts(text):oo.append(part)
    om.append(oo);p._p.append(om)

def para(doc,text):
    p=doc.add_paragraph(text);p.paragraph_format.first_line_indent=Cm(.74);p.paragraph_format.space_after=Pt(5);return p
def heading(doc,text,level=1):doc.add_heading(text,level)
def table(doc,head,rows):
    t=doc.add_table(rows=1,cols=len(head));t.autofit=False
    for j,x in enumerate(head):t.rows[0].cells[j].text=str(x)
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=str(x)
    for i,row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(9);run.bold=(i==0)
            pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for side in ['top','bottom']:
                n=OxmlElement('w:'+side);n.set(qn('w:val'),'single' if i==0 or i==len(t.rows)-1 else 'nil');n.set(qn('w:sz'),'6');n.set(qn('w:color'),'808080');b.append(n)
            pr.append(b)
        rowpr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');rowpr.append(ns)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def build(q):
    folder=m.ROOT/f'问题{q}结果';audit=json.loads((folder/'验证报告.json').read_text(encoding='utf-8'));v=json.loads((m.OUT/'validation.json').read_text(encoding='utf-8'));best=v['best'][str(q)];tot=audit['totals'];checks=audit['checks']
    doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=sec.right_margin=Cm(2.3);sec.top_margin=sec.bottom_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name='Times New Roman';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体');st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(11 if name=='Normal' else (18 if name=='Title' else 13 if name=='Heading 1' else 11))
    doc.styles['Normal'].paragraph_format.line_spacing=1.25
    for name in ['Heading 1','Heading 2']:doc.styles[name].paragraph_format.keep_with_next=True
    title=f'问题{q}微网购电策略建模与求解'
    doc.add_paragraph(title,'Title')
    para(doc,'本章给出经审计后的模型、数据处理、完整公式、因果执行流程和实际计算结果。模型使用附件1的日电价及附件2的全年负荷与光伏；'+('不使用附件3预报。' if q==2 else '同时使用附件3在各发布时刻可获得的光伏预报。')+'所有结果由HiGHS线性规划或混合整数规划求解，按真实储能状态跨日递推。')
    para(doc,f'采用1月15日至31日滚动验证选择的参数，场景数为{best["S"]}，终端价值系数倍率为{best["terminal"]}。参数选择后固定，2月1日至12月31日共334天用于回测。该模型的数值最优性有求解界支持，但未知未来下的全局最佳在线策略不由此得到保证。')
    heading(doc,'1 原模型审计与统一口径')
    common=[('真实执行','整天实际曲线一次优化会提前使用未来信息','每10分钟重解，只有当前实际值进入第一个时段'),('多余供电','弃光不能吸收夜间多购电','另设未利用计划电量z，满足z≤a；弃光w≤v'),('时间轴','输入右端点与模板区间错开10分钟','输入00:10代表00:00至00:10，结果表头纠正'),('储能','充放电损耗不等于互斥约束','LP先给下界；出现同时充放电时转MILP'),('最优性','平衡残差小不能证明费用最优','独立核对费用及原始最优值与求解界'),('初始化','不能在2月1日直接重置6000','1月1日6000起步，整月因果暖机并传递实际SOC')]
    if q==3:common+=[('费用','“退还一半”不是题面唯一解释','主结果采用返还口径，另收违约金口径独立重算'),('误差','相对误差直接减功率存在量纲问题','统一实际减预测的加性电量误差'),('滚动','不调整不等于撤回至0时计划','未更新时延续当前生效版本'),('统计','各天SOC相关，不能随意独立并行','整条策略串行，仅不同策略并行')]
    table(doc,['环节','原稿问题','本版处理'],common)
    para(doc,'已完成逐条模型审查，而不是只改符号。原稿中的“严格多阶段非前视”“误差场景一定提升收益”“损耗保证互斥”等结论均须限定条件。当前采用共同储能轨迹的样本平均计划优化，以及执行层的因果预测控制；不把两阶段完整未来追索当成已经实现的严格多阶段策略。')
    heading(doc,'2 题意约束与数据映射')
    table(doc,['输入','维度与单位','用途'],[['附件1','144个电价，元/kWh','每日重复的固定日电价'],['附件2负荷','365×144，kW','历史预测、执行和验证'],['附件2光伏','365×144，kW','历史误差与执行'],*([['附件3','365×4×24，kW','每6小时发布一次的24小时整点预报']] if q==3 else []),['附件5模板','334日计划及储能、紧急记录','结果表结构']])
    para(doc,'令t=0,…,143，对应[t/6,(t+1)/6)小时。输入的00:10,…,24:00被视为各10分钟区间的右端点功率样本，按矩形积分换算。该采样约定明确记录在时段映射CSV中；题面没有说明小时平均值，不能声称此约定是题面明示。')
    eq(doc,'l_{d,t}=L_{d,t}/6,   v_{d,t}=P_{d,t}/6,   M=5000/6=833.333333')
    para(doc,'模板原表头从00:10至次日00:10，与要求的每日00:00至24:00冲突。结果文件保留工作表名称与144列数量，将表头更正为00:00–00:10至23:50–24:00；不旋转数值、不跨日借用电量。所有四小时汇总、指定时段结果和电价采用同一索引。')
    para(doc,'原始数组均为有限非负数，未发现需填补的缺失记录。负荷和光伏极端值不能使用储能5000kW上限筛除，该上限只约束储能设备。保留真实极值，不做可能抹去风险的标准化或删点；优化直接采用kWh及元。')
    heading(doc,'3 变量定义与核心假设')
    table(doc,['符号','类型 单位','范围与意义'],[['q⁰','连续 kWh','非负，0时合同计划'],['a','连续 kWh','非负，最终生效计划；第二问a=q⁰'],['c r','连续 kWh','均在0至M，分别为充电和放电'],['E','连续 kWh','1200至10800，时段末库存'],['g','连续 kWh','非负，紧急供电'],['w','连续 kWh','0至v，光伏弃电'],['z','连续 kWh','0至a，未利用的计划电量，仍按合同收费'],['u','0或1','u=1充电、u=0放电'],['ηc ηd','已知参数','均为0.9；电池初始量6000kWh'],['λ','校准参数 元/kWh','终端库存价值，1月验证确定']])
    assumptions=[('功率离散化','每10分钟按右端点功率近似平均值','与144段数据对应；短时波动误差可通过更细数据检验'),('无售电收入','多余计划供电可未利用，但费用不退','与按计划计费及供电不低于负荷一致；z不冒充弃光'),('当前信息可观测','当前时段测量可用于本段应急调节','未来时段仅预测；不使用全日实测做控制'),('简化电池','忽略自放电和老化，仅使用题给效率及边界','避免引入无来源参数；不代表实际电池全生命周期模型'),('终端价值','正常日使用线性库存价值，年末至少6000','价值为优化边界处理，不进入题面实际账单')]
    table(doc,['假设','内容','依据与影响'],assumptions)
    heading(doc,'4 预测和场景的可复现构造')
    heading(doc,'4.1 负荷预测',2)
    para(doc,'0时取最近4个同星期日同一时段的加权均值，权重为0.4、0.3、0.2、0.1，历史不足时对已有权重归一化。首周用已有天数均值，1月1日用附件1负荷曲线作初始化先验。')
    eq(doc,'L̂_{d,t}=Σ_{k=1…K} ω_{k} L_{d−7k,t},   Σ_{k=1…K} ω_{k}=1')
    para(doc,'日内控制更新负荷预测：计算过去最多12段（2小时）的实际负荷与0时预测的平均误差，将此偏差以36段为衰减尺度加入未来曲线，并截断到非负。只使用当前时段之前的数据，当前实际值只替换执行优化的第一格。')
    eq(doc,'l̂_{d,j|t}=max(0, l̂_{d,j|0}+b_{d,t} exp(−(j−t)/36))')
    heading(doc,'4.2 光伏预测与时间对齐',2)
    if q==2:
        para(doc,'光伏0时预测取过去7天同一时段的算术均值，1月1日采用零光伏的保守先验。实际执行时，用已观测的最近最多6段光伏误差修正未来预测，衰减尺度为12段。光伏基础预测为零的夜间保持零，不根据未来实测判断日出时刻。')
        eq(doc,'v̂_{d,t|0}=(1/N) Σ_{j=1…N} v_{d−j,t},   N=min(7,d)')
    else:
        para(doc,'附件3第k个预报是发布时刻后第k小时的整点功率。将这些值当作瞬时样本，采用线性插值。发布时刻h的锚点取此前已结束时段的实测；0时锚点取零。对第一个小时，必须在锚点与“预报1小时”之间插值，不能将预报1小时错配到h+1至h+2并漏掉h至h+1。')
        eq(doc,'v̂_{h+j/6}=(1−α) F_{h,k}+α F_{h,k+1},   α=j/6−k')
        para(doc,'插值后先按kW换算为kWh。使用此前最多30天相同发布时刻、相同时段的实际减预测偏差进行加性校准，负值截断为零，夜间零值保持零。日内再以过去最多6段误差做短期衰减修正。全文不混用相对误差和绝对功率误差。')
    heading(doc,'4.3 联合历史误差块',2)
    para(doc,'对每个历史日期，保存当时可得预测的负荷和光伏分量误差，二者均定义为“实际减预测”，单位kWh。选择当前日之前最近60天、且至少已有7天历史的日期作为样本库；负荷与光伏使用同一历史日期，保留日内相关性。')
    eq(doc,'e_{d,t}^{L}=l_{d,t}−l̂_{d,t},   e_{d,t}^{V}=v_{d,t}−v̂_{d,t}')
    eq(doc,'l_{t}^{s}=max(0,l̂_{t}+e_{h_s,t}^{L}−ē_{t}^{L})')
    eq(doc,'v_{t}^{s}=max(0,v̂_{t}+e_{h_s,t}^{V}−ē_{t}^{V})')
    para(doc,'光伏误差仅叠加到基础预测非零的时段。固定随机种子20260911加日期与时段偏移，抽取S个完整日期，样本充足时无放回，不足时有放回，概率均为1/S。误差减去库内均值以免与偏差校准重复。净负荷误差不能唯一恢复两种分量，故本程序从始至终联合保存两者。S=1时直接使用校准点预报，不随机抽一个日块。')
    heading(doc,'5 共同物理约束与最优求解')
    para(doc,'每次优化覆盖当前时刻至日末。计划电量b和参考充放电、储能轨迹在场景间共享，g、w、z可按场景不同。共享储能轨迹是明确的保守计划近似，避免把掌握整个场景路径的储能追索误称为严格在线控制。实际运行由后述每段重解的控制器完成。')
    eq(doc,'b_{t}+v_{t}^{s}−w_{t}^{s}−z_{t}^{s}+r_{t}+g_{t}^{s}=l_{t}^{s}+c_{t}')
    para(doc,'左侧含计划供给、光伏、放电和紧急供给，扣除弃光和未利用计划电量；右侧是负荷与充电需求。若用实际利用的常规外购电x=b−z记号，则等价于x+v−w+r+g=l+c。z不降低账单；w、z分开核查上限，保证来源可追溯。')
    eq(doc,'E_{t}=E_{t−1}+0.9c_{t}−r_{t}/0.9')
    eq(doc,'1200≤E_{t}≤10800,   0≤c_{t}≤M u_{t},   0≤r_{t}≤M(1−u_{t})')
    eq(doc,'u_{t}∈{0,1},   b_{t}≥0,   g_{t}^{s}≥0,   0≤w_{t}^{s}≤v_{t}^{s},   0≤z_{t}^{s}≤b_{t}')
    para(doc,'每次优化的初始库存固定为实测库存。正常日不要求日末等于日初，目标中使用−λE末表示库存价值；年末另外强制E末≥6000，以免全年费用利用初始库存透支。λ=倍率×0.9×日电价中位数，是经1月验证的近似价值，不是题给价格。')
    para(doc,'求解首先移除二进制互斥得到LP下界。若最优LP解已满足互斥，它同时是原MILP可行解，达到下界，因此是原MILP最优解；若出现同时充放电，则显式加入144个或剩余时段数量的0–1变量求解MILP。后者以相对间隙10⁻⁸、60秒时限运行，未达到最优状态则停止，不以启发式结果冒充最优。')
    heading(doc,'6 计划制定与费用推导')
    if q==2:
        para(doc,'第二问每天0时锁定全部144段计划q。对每个场景，紧急供电承担5倍电价。0时求解以下目标，并满足第5节所有约束。各场景的充放电和SOC为共同参考轨迹，计划和轨迹同时优化。')
        eq(doc,'min Σ_{t} p_{t}q_{t}+(1/S)Σ_{s}Σ_{t}5p_{t}g_{t}^{s}−λE_{末}')
        para(doc,'第一项是必须支付的计划费；第二项反映预测误差下的缺电风险；第三项避免把日末库存当作没有价值的能源。实际执行时q不因当天实测值变化，只有储能、弃光、未利用电量和紧急补购更新。')
        eq(doc,'C_{实际}=Σ_{t}p_{t}q_{t}+Σ_{t}5p_{t}g_{t}')
    else:
        para(doc,'0时模型使用第5节约束，并最小化计划费、预测紧急费与终端价值之和。主结果选择S=1时，该问题为校准点预测的确定性MILP。0时计划作为所有调整的合同基准，完整保存。此滚动近似未在0时显式优化未来预报分支，不宣称求得完整多阶段随机规划的全局最优。')
        eq(doc,'min Σ_{t}p_{t}q_{t}^{0}+(1/S)Σ_{s}Σ_{t}5p_{t}g_{t}^{s}−λE_{末}')
        para(doc,'在6时、12时、18时，仅更新未执行区间，当前真实SOC作为新初始状态。决策变量为该版本的调整计划qʰ，不是事后才能确定的最终a。定义δ⁺、δ⁻，满足下式。')
        eq(doc,'q_{t}^{h}−q_{t}^{0}=δ_{t}^{+}−δ_{t}^{−},   δ_{t}^{+}≥0,   δ_{t}^{−}≥0')
        para(doc,'主结果沿用v3中的“下调部分仅收半价，退还另一半”解释，故移除固定的0时计划费用后，调整目标如下。δ⁺和δ⁻同时增加会增加p倍成本，正电价保证最优时至多一个为正，无须额外0–1变量。')
        eq(doc,'min Σ_{t∈R}p_{t}(−0.5δ_{t}^{−}+1.5δ_{t}^{+})+(1/S)Σ_{s}Σ_{t∈R}5p_{t}g_{t}^{s}−λE_{末}')
        para(doc,'原始题面“违约电价50%”未明确退款。另一个可辩护的解释是合同计划全价照付，再对下调加收0.5倍违约金。该口径须将上述−0.5改为+0.5，并重新优化，不能只对同一策略改账单。包内提供这一独立回测的对照结果，正文不把其中任一解释宣称为题面已明确事实。')
        table(doc,['合同q⁰','调整a','退款口径费用','另收违约金费用'],[['100','80','90p','110p'],['100','120','130p','130p']])
        eq(doc,'C=Σ_{t}p_{t}(q_{t}^{0}−0.5δ_{t}^{−}+1.5δ_{t}^{+}+5g_{t})')
        eq(doc,'C=Σ_{t}p_{t}(a_{t}+0.5|a_{t}−q_{t}^{0}|+5g_{t})')
        para(doc,'两式在退款口径下严格相等，第二式可用作独立核算。不得在第一式中将q⁰换成a又保留相同调整项。额外收费口径不满足第二个等价式。')
        heading(doc,'6.1 滚动版本的拼接',2)
        table(doc,['区间','最终生效计划'],[['00:00–06:00','0时版本'],['06:00–12:00','6时版本，如未采用则沿用此前版本'],['12:00–18:00','12时版本，如未采用则沿用当前版本'],['18:00–24:00','18时版本，如未采用则沿用当前版本']])
        para(doc,'6时计划可覆盖至日末，但只在下一发布之前被执行；12时到达后，过去执行量冻结，未来由新版本覆盖。若不引入某发布时间，当前版本自然继续生效，而不是强制退回0时计划。版本拼接使用保存的144段数组独立验证。')
    heading(doc,'7 实际执行和跨日回测')
    para(doc,'在每个10分钟时段t，读取当前负荷、光伏与起始库存；余下时段用当前可得预测。第二问锁定q，第三问锁定截至此时有效的a。以当前至日末的预测序列求解以下确定性控制问题，仅执行第一个时段的c、r、g、w、z，随后按真实效率递推SOC。下一段重新求解，决不执行“全日真实数据已知”的离线轨迹。')
    eq(doc,'min Σ_{j=t…143}5p_{j}g_{j}−λE_{143},   b_{j}=a_{j} 已知')
    para(doc,'所有物理、来源和互斥约束仍保留。计划费用在此阶段已经固定，移除不会改变最优控制。实际账单只累计已执行第一段的紧急购电费用，不计未发生的未来预测费用，也不计终端价值。当前分配的来源变量可能存在多解，程序在同等费用和储能轨迹下优先把剩余供给归入弃光，再归入未利用计划电量；因此w、z分解是核算约定，总供给剩余与总费用不变。')
    para(doc,'完整流程：检查原始数据→1月1日6000kWh初始化→每日0时预测与计划优化→每段观测和预测控制→第三问在发布点更新计划→累计实际账单→跨日传递实际SOC→2月至12月结果导出→独立验证。每条策略的365天内部串行，不固定每天初始SOC来伪造独立性。')
    heading(doc,'8 参数选择与方案比较')
    cand=[r for r in v['candidates'] if r['q']==q]
    table(doc,['场景数','终端倍率','1月15至31日费用 元'],[[r['S'],r['terminal'],f'{r["cost"]:,.2f}'] for r in cand])
    para(doc,'这里的“较优”严格指上述候选在1月验证集上的最低实际费用；不利用2月至12月挑选参数。场景数比较用于检查复杂化的收益，不将“取50或100场景”当作已完成收敛证明。终端倍率比较同时覆盖库存价值敏感性。早期历史样本较短，验证选择可能有方差，结果的泛化能力只能由随后固定策略的回测评估。')
    if q==3:
        comp=json.loads((folder/'比较摘要.json').read_text(encoding='utf-8'))
        table(doc,['策略','334天总费用 元','紧急电量 kWh'],[[r[0],f'{r[1]:,.2f}',f'{r[2]:,.2f}'] for r in comp['rows']])
        table(doc,['新增信息','日均节约 元','95%区间'],[[r[0],f'{r[2]:,.2f}',f'[{r[3]:,.2f}, {r[4]:,.2f}]'] for r in comp['cis']])
        para(doc,'各策略从相同的1月1日6000kWh出发，随后保留各自的SOC轨迹。置信区间采用7天移动块Bootstrap，重复2000次，保留一周内相关性。季节和光伏出力三分位分层仅用于结果解释，不把低出力日直接判作阴天。区间包含0表示当前证据不足以断言新增信息稳定节约费用。')
        para(doc,'本版未启用事后统计显著性作为在线开关。若扩展收益门槛，应在同一当前信息集和同一场景集下分别计算“保留当前计划”与“重新优化”的预测费用，并比较Δ=J保留−J更新与事先指定阈值。零阈值在精确优化时通常退化为总是允许更新，不能宣称它自动过滤噪声。')
    heading(doc,'9 实际结果与完整性验证')
    table(doc,['费用与电量指标','数值'],[['计划费用 元',f'{tot["plan_cost"]:,.2f}'],['调整费用 元',f'{tot["adjust_cost"]:,.2f}'],['紧急费用 元',f'{tot["emergency_cost"]:,.2f}'],['总费用 元',f'{tot["total_cost"]:,.2f}'],['紧急电量 kWh',f'{tot["emergency_kWh"]:,.2f}'],['2月1日初始SOC kWh',f'{checks["feb_start_SOC"]:,.3f}'],['年末SOC kWh',f'{checks["year_end_SOC"]:,.3f}']])
    table(doc,['独立检查','结果'],[['最大能量平衡残差 kWh',f'{checks["energy_balance_max"]:.3g}'],['最大储能递推残差 kWh',f'{checks["soc_recursion_max"]:.3g}'],['最大费用重算偏差 元',f'{checks["cost_reconciliation_max"]:.3g}'],['最小 最大SOC kWh',f'{checks["soc_min"]:.3f} / {checks["soc_max"]:.3f}'],['最大充放电重叠 kWh',f'{checks["charge_discharge_overlap"]:.3g}'],['弃光来源越界 kWh',f'{checks["pv_curtail_upper_violation"]:.3g}'],['最大最优界差 元',f'{checks["absolute_optimality_gap_max_CNY"]:.3g}'],['优化次数 含MILP次数',f'{checks["lp_or_mip_solves"]} / {checks["mip_solves"]}']])
    para(doc,'针对没有光伏却计划多购、计划不足需紧急供电、终端补库效率、两种偏差结算手算例以及未来数据扰动的测试均由代码执行。求解日志记录每次目标值、下界和间隙，不能用图形平滑或平衡残差替代最优性证据。')
    doc.add_picture(str(folder/'指定日期运行图.png'),width=Cm(16));para(doc,'图1 四个指定日期的实际运行与储能轨迹。电量按10分钟汇总，SOC按真实充放电递推。')
    if q==3:doc.add_picture(str(folder/'信息价值比较.png'),width=Cm(16));para(doc,'图2 固定参数与相同结算口径下的预报策略比较。')
    heading(doc,'10 指定日期与交付文件')
    para(doc,'论文指定日期为2025年3月20日、6月21日、9月23日、12月21日。结果包中的指定日期表1.csv列出10、12、14、16、18、20时起的10分钟购电量，以及全天电量和分项费用；表2.csv列出六个四小时区间的充放电与首末SOC；表3.csv列出连续紧急购电区间，无紧急购电则明确填0。')
    para(doc,f'result{q}.xlsx的计划购电量表包含334天、144段与全天计划电量及费用。'+('调整购电量表保存拼接后的最终生效计划a，全天费用列为包含计划、调整与紧急的实际账单。' if q==3 else '')+'充放电量表完整展开每天六行，将0时SOC放在第一行、24时SOC放在第二行，避免标签和数值错行。紧急记录按同日相邻正g时段合并，24时明确写作24:00。')
    para(doc,'结果包还包含逐时段审计数据、逐日费用、时段映射、验证报告、求解最优性日志、参数验证表和可重跑代码。模型与代码逐项一致；未把可选场景树、K-medoids或天气模型写作已实现功能。')
    heading(doc,'11 局限与下一步改进边界')
    para(doc,'当前计划层的共同储能轨迹可能偏保守；执行层未来采用点预测而非完整多阶段场景树；线性库存价值只是跨日价值的近似；1月短验证集不覆盖全年季节；整点插值与输入区间语义仍依赖明确采样假设。以上局限应在论文中保留。下一步若扩展场景树，必须在相同历史节点共享储能及购电动作，再比较计算成本与样本外收益。')
    para(doc,'本次“最优”指给定输入、已声明模型及数值容差下的各次优化最优，和候选参数在1月验证集上的最优；不等价于真实未来已知时的最低费用，不构成任何奖项承诺。')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    out=m.ROOT/f'第{q}问建模文档_核验重算版.docx';doc.save(out);print(out,flush=True)

if __name__=='__main__':
    build(2);build(3)
