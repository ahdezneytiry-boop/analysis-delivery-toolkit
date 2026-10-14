"""两问共用的真实线性优化与因果回测。Python >=3.12，numpy scipy openpyxl。
修改 INPUT 为附件目录；OUT 为独立输出目录。所有历史切片严格为 [:day]。
计划采用共同储能轨迹的样本平均优化；执行采用每10分钟重解的确定性MPC。
禁止把本程序的模型最优性解释为未知未来下的全局最优策略。
"""
from pathlib import Path
import sys, json, time, hashlib, argparse
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'deps'))
import numpy as np
from scipy.optimize import linprog, milp, Bounds, LinearConstraint
from scipy.sparse import lil_matrix, csr_matrix, vstack, hstack, eye
from functools import lru_cache
from openpyxl import load_workbook

INPUT=Path(r'C:/Users/ahdezneytiry-boop1/Desktop/附件') # <<< 修改为本地附件目录
OUT=ROOT/'计算结果' # <<< 所有结果输出在这里
ETA=.9; M=5000/6; EMIN=1200.; EMAX=10800.; INITIAL=6000.
SOLVES=[]

@lru_cache(maxsize=1000)
def matrices(S,T,adjust):
    n=4*T+3*S*T+(2*T if adjust else 0)
    rr=[];cc=[];vv=[]; ur=[];uc=[];uv=[]
    for s in range(S):
        for t in range(T):
            row=s*T+t; g=4*T+s*T+t
            rr.extend([row]*6);cc.extend([t,T+t,2*T+t,g,g+S*T,g+2*S*T]);vv.extend([1,-1,1,1,-1,-1])
            ur.extend([row]*2);uc.extend([g+S*T,t]);uv.extend([1,-1])
    for t in range(T):
        row=S*T+t; rr.extend([row]*3);cc.extend([3*T+t,T+t,2*T+t]);vv.extend([1,-ETA,1/ETA])
        if t:rr.append(row);cc.append(3*T+t-1);vv.append(-1)
    if adjust:
        for t in range(T):
            row=S*T+T+t;rr.extend([row]*3);cc.extend([t,4*T+3*S*T+t,4*T+3*S*T+T+t]);vv.extend([1,-1,1])
    return csr_matrix((vv,(rr,cc)),shape=(S*T+T+(T if adjust else 0),n)),csr_matrix((uv,(ur,uc)),shape=(S*T,n))

def data():
    w=load_workbook(INPUT/'附件2.xlsx',data_only=True,read_only=True)
    rows=list(w.worksheets[0].values); dates=[r[0].strftime('%Y-%m-%d') for r in rows[1:]]
    L=np.array([r[1:] for r in rows[1:]],float)/6
    V=np.array([r[1:] for r in list(w.worksheets[1].values)[1:]],float)/6
    w=load_workbook(INPUT/'附件1.xlsx',data_only=True,read_only=True)
    r=list(w.active.values)[1:]; p=np.array([x[1] for x in r]); base=np.array([x[2] for x in r])/6
    r=list(load_workbook(INPUT/'附件3.xlsx',data_only=True,read_only=True).active.values)[1:]
    F=np.array([x[2:] for x in r],float).reshape(365,4,24)/6
    assert L.shape==V.shape==(365,144) and np.isfinite(L).all() and np.isfinite(V).all()
    assert L.min()>=0 and V.min()>=0 and p.min()>0
    return dates,L,V,p,base,F

class Forecast:
    def __init__(self,L,V,base,F):
        self.L=L; self.V=V; self.F=F; self.l=[]; self.v=[]; self.raw=np.zeros((365,4,144))
        for d in range(365):
            ids=[d-7*k for k in range(1,5) if d-7*k>=0]
            if ids:
                w=np.array([.4,.3,.2,.1][:len(ids)]); w/=w.sum(); lf=w@L[ids]
            else: lf=L[max(0,d-7):d].mean(0) if d else base.copy()
            vf=V[max(0,d-7):d].mean(0) if d else np.zeros(144)
            self.l.append(lf); self.v.append(vf)
            for j,h in enumerate([0,36,72,108]):
                # 整点为瞬时功率，右端点10分钟采样：发布时刻实测仅取已结束时段。
                anchor=V[d,h-1] if h else 0.
                knots=np.r_[anchor,F[d,j]]
                self.raw[d,j,h:]=np.interp(np.arange(h+1,145),h+np.arange(25)*6,knots)
        self.l=np.array(self.l); self.v=np.array(self.v)

    def predict(self,d,t,q,release=0):
        lf=self.l[d].copy()
        # 仅使用当前时段之前的负荷观测，近期误差指数衰减。
        if t:
            bias=np.mean(self.L[d,max(0,t-12):t]-lf[max(0,t-12):t])
            lf[t:]+=bias*np.exp(-np.arange(144-t)/36)
        if q==2: vf=self.v[d].copy()
        else:
            j=release//36; vf=self.raw[d,j].copy()
            ids=np.arange(max(0,d-30),d)
            if len(ids):
                # 误差实际-预测，单位kWh，校准相加；夜间保持0。
                bias=np.mean(self.V[ids]-self.raw[ids,j],axis=0)
                vf+=bias*(vf>1e-8)
        if t:
            bias=np.mean(self.V[d,max(release,t-6):t]-vf[max(release,t-6):t]) if t>release else 0.
            vf[t:]+=bias*np.exp(-np.arange(144-t)/12)*(vf[t:]>1e-8)
        return np.maximum(lf,0),np.maximum(vf,0)

    def scenarios(self,d,t,q,release,S):
        lf,vf=self.predict(d,t,q,release)
        ids=np.arange(max(7,d-60),d)
        if S==1 or not len(ids):return lf[None,t:],vf[None,t:]
        # 固定种子，对同日负荷和光伏联合抽样，不逐时段独立抽样。
        rng=np.random.default_rng(20260911+d*144+t)
        ix=rng.choice(ids,size=S,replace=len(ids)<S)
        el=self.L[ix]-self.l[ix]
        ev=self.V[ix]-(self.v[ix] if q==2 else self.raw[ix,release//36])
        # 历史误差中心化，预测偏差已在predict内校准。
        el-=np.mean(self.L[ids]-self.l[ids],axis=0)
        ev-=np.mean(self.V[ids]-(self.v[ids] if q==2 else self.raw[ids,release//36]),axis=0)
        return np.maximum(0,lf[None,t:]+el[:,t:]),np.maximum(0,vf[None,t:]+ev[:,t:]*(vf[None,t:]>1e-8))

def opt(L,V,p,e0,base=None,fixed=None,fee=-.5,terminal=0.,target=None):
    """LP变量：b,c,r,E(共同), g_s,z_s,w_s；z=已付但未消纳，w=弃光。
    b-z+V-w+r+g=L+c。z<=b，w<=V。b为计划/调整量或固定执行量。
    base非空时目标按 q0 + fee*delta- +1.5delta+，固定项移除。
    先求LP下界；若LP解不互斥则加入二进制变量求MILP，保留最优界。
    """
    S,T=L.shape; n=4*T+3*S*T+(2*T if base is not None else 0)
    B=np.arange(T); C=B+T; R=B+2*T; E=B+3*T
    G=np.arange(S*T).reshape(S,T)+4*T; Z=G+S*T; W=G+2*S*T
    f=np.zeros(n); bounds=[(0,None)]*n
    if base is None:f[B]=p
    if fixed is not None:
        for t in range(T):bounds[B[t]]=(float(fixed[t]),float(fixed[t]))
        f[B]=0
    for t in range(T):bounds[C[t]]=(0,M);bounds[R[t]]=(0,M);bounds[E[t]]=(EMIN,EMAX)
    f[G.ravel()]=np.tile(5*p/S,S);f[E[-1]]=-terminal
    if target is not None:bounds[E[-1]]=(target,EMAX)
    for s in range(S):
        for t in range(T):bounds[W[s,t]]=(0,float(V[s,t]))
    A,U=matrices(S,T,base is not None);rhs=np.r_[(L-V).ravel(),e0,np.zeros(T-1),base if base is not None else []]
    if base is not None:
        DP=np.arange(T)+4*T+3*S*T;DM=DP+T
        f[DP]=1.5*p;f[DM]=fee*p
    res=linprog(f,A_ub=U.tocsr(),b_ub=np.zeros(S*T),A_eq=A.tocsr(),b_eq=rhs,bounds=bounds,method='highs')
    if not res.success:raise RuntimeError(res.message)
    lpbound=res.fun; was_mip=False
    if np.max(np.minimum(res.x[C],res.x[R]))>1e-6:
        # 松弛不保证互斥；显式二进制MILP精解，不用损耗当作互斥证明。
        H=lil_matrix((2*T,n+T))
        for t in range(T):H[t,[C[t],n+t]]=[1,-M];H[T+t,[R[t],n+t]]=[1,M]
        zero=csr_matrix((A.shape[0],T));uzero=csr_matrix((U.shape[0],T))
        constraints=[LinearConstraint(hstack([A,zero]).tocsr(),rhs,rhs),LinearConstraint(hstack([U,uzero]).tocsr(),-np.inf,np.zeros(S*T)),LinearConstraint(H.tocsr(),-np.inf,np.r_[np.zeros(T),np.full(T,M)])]
        lo=np.array([0 if x[0] is None else x[0] for x in bounds]+[0]*T);hi=np.array([np.inf if x[1] is None else x[1] for x in bounds]+[1]*T)
        res=milp(np.r_[f,np.zeros(T)],integrality=np.r_[np.zeros(n),np.ones(T)],bounds=Bounds(lo,hi),constraints=constraints,options={'mip_rel_gap':1e-8,'time_limit':60})
        if not res.success:raise RuntimeError('MILP未证明最优: '+res.message)
        was_mip=True
    x=res.x; b=x[B];c=x[C].copy();r=x[R].copy(); e=x[E].copy()
    # 从最优充放电重新构造最少g，再将多余电量分别归入w或z。
    k=np.zeros(T)
    g=np.maximum(L+c[None,:]-V-b[None,:]-r[None,:],0)
    surplus=np.maximum(b[None,:]+V+r[None,:]-L-c[None,:],0)
    w=np.minimum(V,surplus);z=surplus-w
    bal=b+V-w-z+r+g-L-c
    assert np.max(abs(bal))<1e-6 and np.max(z-b)<1e-6, (float(np.max(abs(bal))),float(np.max(z-b)),float(k.max()),T,L.min(),res.fun)
    assert np.max(abs(e-np.cumsum(ETA*c-r/ETA)-e0))<1e-5
    assert np.max(np.minimum(c,r))<1e-6
    obj=(0 if fixed is not None else np.dot(p,b)) if base is None else np.dot(p,1.5*np.maximum(b-base,0)+fee*np.maximum(base-b,0))
    obj+=np.mean(g@(5*p))-terminal*e[-1]
    assert abs(obj-res.fun)<1e-4, (obj,res.fun)
    # HiGHS返回最优状态，原始对偶值独立核对。
    dual=float(res.mip_dual_bound) if was_mip else float(rhs@res.eqlin.marginals)
    if not was_mip:
        for i,(lo,hi) in enumerate(bounds):
            if lo is not None:dual+=lo*res.lower.marginals[i]
            if hi is not None:dual+=hi*res.upper.marginals[i]
    gap=abs(res.fun-dual)
    SOLVES.append((T,S,float(res.fun),dual,gap,int(was_mip)))
    return dict(b=b,c=c,r=r,E=e,g=g,w=w,z=z,obj=float(obj),gap=gap)

def run(q,S=10,updates=(36,72,108),fee=-.5,end=365,begin=0,terminal_scale=1.,save=True,name=None):
    dates,L,V,p,base,F=data(); fc=Forecast(L,V,base,F); e0=INITIAL
    records=[]; daily=[]; plans=[]; versions=[]; releases=[]
    # Jan warm-up is same causal policy; never set Feb 1 directly to6000.
    starttime=time.time(); solve_start=len(SOLVES)
    for d in range(begin,end):
        sl,sv=fc.scenarios(d,0,q,0,S)
        term=terminal_scale*ETA*float(np.median(p)); target=INITIAL if d==364 else None
        sol=opt(sl,sv,p,e0,terminal=term,target=target); q0=sol['b'].copy(); cur=q0.copy(); release=0
        day=[]; ver=np.full((4,144),np.nan);ver[0]=q0.copy(); ebegin=e0
        for t in range(144):
            if q==3 and t in updates:
                release=t
                sl,sv=fc.scenarios(d,t,q,t,S)
                adj=opt(sl,sv,p[t:],e0,base=q0[t:],fee=fee,terminal=term,target=target)
                cur[t:]=adj['b'];ver[t//36,t:]=cur[t:]
            # 实际执行仅当前时段L,V已揭示，未来用过去信息预测，每段重解。
            lf,vf=fc.predict(d,t,q,release)
            lf[t]=L[d,t];vf[t]=V[d,t]
            ex=opt(lf[None,t:],vf[None,t:],p[t:],e0,fixed=cur[t:],terminal=term,target=target)
            c,r,g,w,z=float(ex['c'][0]),float(ex['r'][0]),float(ex['g'][0,0]),float(ex['w'][0,0]),float(ex['z'][0,0])
            en=float(ex['E'][0]); a=float(cur[t]);down=max(q0[t]-a,0);up=max(a-q0[t],0)
            adjcost=p[t]*(fee*down+1.5*up) if q==3 else 0.
            day.append([d,t,L[d,t],V[d,t],p[t],q0[t],a,c,r,g,w,z,e0,en,down,up,p[t]*q0[t],adjcost,5*p[t]*g])
            e0=en
        day=np.array(day);records.append(day);plans.append(q0);versions.append(ver)
        daily.append([d,day[:,16].sum(),day[:,17].sum(),day[:,18].sum(),day[:,6].sum(),day[:,9].sum(),day[:,10].sum(),day[:,11].sum(),ebegin,e0])
        if (d+1)%7==0 or d==end-1:print(json.dumps({'run':name or f'q{q}_S{S}','day':d+1,'sec':round(time.time()-starttime,1),'cost':round(sum(day[:,16:19].sum(axis=0)),2)},ensure_ascii=False),flush=True)
    result=dict(records=np.array(records),daily=np.array(daily),plans=np.array(plans),versions=np.array(versions),dates=np.array(dates[begin:end]),solves=np.array(SOLVES[solve_start:]))
    if save:
        OUT.mkdir(parents=True,exist_ok=True);path=OUT/(name or f'q{q}_S{S}');np.savez_compressed(path,**result)
    return result

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--q',type=int,default=2);ap.add_argument('--S',type=int,default=10);ap.add_argument('--end',type=int,default=365);ap.add_argument('--updates',default='36,72,108');ap.add_argument('--fee',type=float,default=-.5);ap.add_argument('--name');ap.add_argument('--terminal',type=float,default=1.)
    a=ap.parse_args();run(a.q,a.S,tuple(int(i) for i in a.updates.split(',') if i),a.fee,a.end,terminal_scale=a.terminal,name=a.name)
