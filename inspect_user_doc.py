from docx import Document
from pathlib import Path
p=Path('C:/Users/ahdezneytiry-boop1/Doubao/chats/2026-09-11/new-chat/modeling_doc/C题建模文档_北大核心格式.docx')
d=Document(p)
for i,p in enumerate(d.paragraphs):
 print(i,p.style.name,repr(p.text))
for i,t in enumerate(d.tables):
 print('TABLE',i,len(t.rows),len(t.columns))
 for r in t.rows:print([c.text for c in r.cells])
print('SECTIONS',[(s.page_width.cm,s.page_height.cm,s.top_margin.cm,s.bottom_margin.cm,s.left_margin.cm,s.right_margin.cm) for s in d.sections])
print('IMAGES',len(d.inline_shapes))
