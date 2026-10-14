from pathlib import Path
import sys,json,hashlib,zipfile
ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'python_deps'))
import pandas as pd
import numpy as np
import openpyxl
from scipy.io import loadmat

data=ROOT/'预处理数据';result=ROOT/'第一问结果'
a=pd.read_csv(data/'actual_10min.csv');f=pd.read_csv(data/'pv_forecast_hourly.csv');q=pd.read_csv(data/'q1_clean.csv')
d=pd.read_csv(result/'q1_dispatch.csv');m=loadmat(data/'clean_data.mat')
assert len(a)==52560 and len(f)==35040 and len(q)==144
for name,col in [('load_kW','load_kW'),('pv_actual_kW','pv_actual_kW'),('price','price_yuan_per_kWh')]:
    assert np.allclose(m[name].ravel(),a[col],rtol=0,atol=1e-7)
assert np.allclose(a.load_kW/6,a.load_kWh,atol=1e-6)
assert np.allclose(a.pv_actual_kW/6,a.pv_actual_kWh,atol=1e-6)
assert np.allclose(m['forecast_pv_kW'].ravel(),f.pv_forecast_kW,atol=1e-7)
assert d.soc_end_kWh.between(1200-1e-6,10800+1e-6).all()
assert d.charge_ac_kWh.between(-1e-6,5000/6+1e-6).all()
assert d.discharge_ac_kWh.between(-1e-6,5000/6+1e-6).all()
assert (d.curtailment_kWh<=d.pv_forecast_kWh+1e-6).all()
wb=openpyxl.load_workbook(result/'result1.xlsx',data_only=True)
assert wb.sheetnames==['计划购电量','充放电量']
sheet=wb['计划购电量'];values=np.array([sheet.cell(i,2).value for i in range(2,146)])
assert np.allclose(values,d.purchase_kWh,atol=1e-6)
assert sheet['A2'].value=='00:00-00:10' and sheet['A145'].value=='23:50-24:00'
assert abs(sheet['E2'].value-values.sum())<1e-5
assert abs(sheet['E3'].value-float(d.cost_yuan.sum()))<1e-5
items=[ROOT/'全题解析与第一问建模.md',ROOT/'preprocess_and_q1.py']
items+=list(data.glob('*'))
items+=[result/n for n in ['result1.xlsx','q1_dispatch.csv','q1_blocks.csv','q1_sensitivity.csv','verification.json','q1_dispatch.png']]
lines=['# 首轮交付文件清单','已校验CSV/MAT数值一致性、时间键、能量单位、储能边界及XLSX与求解结果一致。','', '|文件|大小/字节|格式|SHA-256|','|---|---:|---|---|']
for p in items:
    lines.append(f'|{p.relative_to(ROOT).as_posix()}|{p.stat().st_size}|{p.suffix[1:]}|{hashlib.sha256(p.read_bytes()).hexdigest()}|')
manifest=ROOT/'文件清单.md';manifest.write_text('\n'.join(lines),encoding='utf-8');items.append(manifest)
zip_path=ROOT/'C题_首轮全题框架与第一问.zip'
with zipfile.ZipFile(zip_path,'w',compression=zipfile.ZIP_DEFLATED) as z:
    for p in items:z.write(p,p.relative_to(ROOT))
print('Integrity checks passed.')
print('ZIP',zip_path.stat().st_size)
for p in [ROOT/'全题解析与第一问建模.md',ROOT/'preprocess_and_q1.py',data/'q1_clean.csv',data/'actual_10min.csv',data/'pv_forecast_hourly.csv',data/'clean_data.mat',result/'result1.xlsx']:print(p.name,p.stat().st_size)
