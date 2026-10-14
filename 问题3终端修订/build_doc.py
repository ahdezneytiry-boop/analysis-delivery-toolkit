"""生成完整第三问文档，不将敏感性结果只追加为孤立附件。"""
import base_doc as b
import solve as m,json
from docx import Document
from docx.shared import Cm
def section8(doc):
    s=json.loads((m.ROOT/'问题3结果'/'汇总.json').read_text(encoding='utf8'));selected=s['selection']['selected_alpha'];rows=s['rows'];c=s['comparison']
    b.heading(doc,'8 终端价值标定与阈值检验')
    b.para(doc,'终端价值是有限日内预测窗口的跨日近似，不是题给电价或物理常数。采用无量纲α表示库存偏好，λ为元/kWh。代码以alpha为主入口，旧接口terminal_scale仅用于兼容，关系如下。')
    b.eq(doc,'λ_{d}=α median(p)  (d<365),   λ_{365}=0,   terminal_scale=α/0.9')
    b.para(doc,'所有候选固定S=1、相同预报校准、相同结算及执行规则。候选为0.30、0.50、0.51、0.52、0.55、0.60、0.65、0.70、0.90、1.10。各候选从1月1日6000kWh因果运行，仅按1月15至31日现金费用排序，之后参数固定。2月至12月费用单列，不能反向决定所选参数。')
    b.para(doc,'每个候选的1月轨迹是策略评估中的暖机路径，2月初SOC由该路径继承，并不表示实际电池能在选参后回到另一个历史状态。实际部署必须用选参时真实观测SOC重新起算；本文报告的是统一1月初状态的完整策略回测。')
    b.eq(doc,'α* = argmin_{α∈A} Σ_{d=15…31} C_{d}(α;S=1)')
    b.para(doc,f'本次历史验证选得α={selected:.2f}。候选集合是在审查既有全年结果后扩充的，因此本版属于事后模型修订，不能宣称得到完全未经开发接触的独立测试集。相较旧表，本版消除了场景数与终端系数混合比较的问题；仍不把有限候选最优称为连续参数全局最优。')
    b.table(doc,['α','1月验证 元','2—12月 元','日末SOC均值','满电比例'],[[f'{r[0]:.2f}',f'{r[2]:,.2f}',f'{r[3]:,.2f}',f'{r[6]:.1f}',f'{100*r[5]:.2f}%'] for r in rows])
    retrospective=min(rows,key=lambda r:r[3]);chosen=next(r for r in rows if r[0]==selected)
    b.para(doc,f'在本次候选中，2—12月事后最低费用对应α={retrospective[0]:.2f}，费用{retrospective[3]:,.2f}元；历史验证主参数α={selected:.2f}的全年费用比它高{chosen[3]-retrospective[3]:,.2f}元。二者分别回答“已知全年后哪项最低”和“按既定历史验证规则选择哪项”，必须同时报告，不以信息价值方向稳定掩盖参数费用差。')
    b.para(doc,'满电比例定义为时段末E≥10799.999kWh的时段占比；日末SOC均值只取每日24时库存。时段SOC均值使用144个时段末状态等权平均，避免把某一天初值或末值误称为全年均值。')
    b.table(doc,['α','弃光 kWh','未利用购电 kWh','年末SOC kWh'],[[f'{r[0]:.2f}',f'{r[7]:,.1f}',f'{r[8]:,.1f}',f'{r[10]:,.1f}'] for r in rows])
    b.para(doc,'低谷电价为0.3713元/kWh，电价中位数为0.80735元/kWh。若边际库存完全通过低谷计划电充入，则增加1kWh库存的成本为p谷/ηc，与终端奖励λ比较，可得候选切换系数约0.511。它只解释一种边际情形；已锁定合同、调整价格、免费剩余光伏和功率约束可能产生其他切换点。')
    b.eq(doc,'α_{候选}=p_{谷}/(η_{c} median(p))≈0.511')
    b.para(doc,'因此在0.50至0.52及0.55至0.65补充取值。线性规划的活跃约束切换和整数互斥可能产生跳变，不要求结果平滑；当前离散网格只能定位变化区间，不能宣称识别唯一精确阈值。逐日敏感性文件保留全部日期的费用、日末SOC、满电比例、弃光和未利用购电量。')
    tr=s['largest_daily_SOC_change_interval'];b.para(doc,f'在本次网格中，日末SOC均值变化最大的相邻区间为α={tr[0]:.2f}至{tr[1]:.2f}，变化{tr[3]:.2f}kWh，相应总费用变化{tr[2]:.2f}元。该区间是后续局部细分的候选范围，尚不证明区间内只有一个切换点。')
    doc.add_picture(str(m.ROOT/'问题3结果'/'终端系数敏感性.png'),width=Cm(16))
    b.para(doc,'图1 统一年末边界后的费用与库存敏感性。连线仅帮助辨认采样点，不表示中间系数已计算。')
    b.heading(doc,'8.1 年末边界与公平评价',2)
    b.para(doc,'正常日仍以−λdE日末近似次日价值。12月31日的0时计划、每次调整和每段执行全部使用λ365=0，同时保留E365,143≥6000的期末库存约束。这是本研究的评价边界假设，不是题面明示要求；不要求每日日末等于日初。若免费剩余供电导致库存大于6000，允许保留，但不在现金账单中计入虚构收益。')
    b.para(doc,'原有五个系数的前364天规则未改变，复用其逐段轨迹和求解日志，重新计算12月31日。新增系数全年从头执行。程序拼接时核验12月30日末与12月31日初SOC完全连续；复现脚本也可全部从1月1日重跑，不依赖缓存。')
    b.heading(doc,'8.2 信息价值与第二问的衔接',2)
    b.para(doc,'本次信息价值比较仅改变是否使用6、12、18时预报与调整：两策略固定同一个α、S=1、负荷预测规则、物理边界和退款结算。仅0时策略仍使用附件3的0时预报，不等同于第二问的历史光伏预测策略。两策略从1月1日相同库存起步，2月初库存允许因真实暖机而不同。')
    b.table(doc,['策略','2—12月费用 元'],[['仅0时预报',f'{c["none_cost"]:,.2f}'],['全部预报',f'{c["all_cost"]:,.2f}'],['差额 仅0时减全部',f'{c["saving"]:,.2f}']])
    b.para(doc,f'采用7日移动块Bootstrap重复2000次，日均节约的95%区间为[{c["daily_saving_95CI"][0]:.2f}, {c["daily_saving_95CI"][1]:.2f}]元。该区间仅描述这组策略的序列波动，未消除候选集合事后修订的不确定性。')
    b.table(doc,['α','仅0时费用 元','全部预报费用 元','节约 元'],[[f'{r[0]:.2f}',f'{r[1]:,.2f}',f'{r[2]:,.2f}',f'{r[3]:,.2f}'] for r in s['paired_information_value']])
    b.table(doc,['项目','第二问既有版','第三问本版'],[['物理约束','相同储能边界和效率','相同储能边界和效率'],['光伏信息','历史滚动预测','附件3分次预报'],['合同调整','日内锁定','发布点调整'],['场景数','10','固定1'],['终端系数α','0.45',f'{selected:.2f}'],['年末处理','原版仍有终端价值','末日价值为0且库存≥6000']])
    b.para(doc,'因算法参数和年末规则仍不同，不能用既有第二问总费用直接证明第三问的纯信息价值。本版不重算第二问，也不将两问差额全部归因于预报。若需要严格跨题消融，应统一上述算法配置和评价边界后另行计算。')
    paired=s['paired_information_value'];direction='全部预报费用均更低' if all(r[3]>0 for r in paired) else '出现节约方向不一致'
    b.para(doc,f'原0.9配置并非已测试配置中的最低费用。本次统一年末边界的成对比较中，{direction}。信息价值的方向与绝对费用的参数敏感性是两项不同结论；不能以前者掩盖后者。')

def main():
    s=json.loads((m.ROOT/'问题3结果'/'汇总.json').read_text(encoding='utf8'));a=s['selection']['selected_alpha'];cand=[{'S':1,'terminal':r['alpha']/.9,'cost':r['jan_validation']} for r in s['selection']['candidates']]
    (m.OUT/'validation.json').write_text(json.dumps({'candidates':[dict(r,q=3) for r in cand],'best':{'3':{'S':1,'terminal':a/.9}}}),encoding='utf8')
    b.section8=section8;b.build(3)
    path=m.ROOT/'第3问建模文档_核验重算版.docx';doc=Document(path)
    for p in doc.paragraphs:
        if p.text.startswith('图1 四个指定日期'):p.text=p.text.replace('图1','图2',1)
        if p.style.name=='Title':p.text='第三问微网购电策略与终端价值修订';ppr=p._p.get_or_add_pPr();border=ppr.find(b.qn('w:pBdr'));ppr.remove(border) if border is not None else None
        if p.text.startswith('采用1月15日至31日滚动验证选择的参数'):
            p.text=f'本版固定场景数S=1，对10个终端系数进行比较，按1月15日至31日现金费用选得α={a:.2f}，正常日λ=α×median(p)。12月31日λ置零，库存不低于6000kWh。2月至12月结果作为修订模型回测报告，候选集合的事后扩充已明确披露。'
        if p.text.startswith('每次优化的初始库存固定为实测库存'):
            p.text='每次优化从真实库存出发，正常日不要求日末等于日初。目标中−λdE末只表示跨日近似价值。正常日λd=α×median(p)，最后一天λ365=0且E末≥6000。所有终端价值均不进入实际现金账单。'
        if p.text.startswith('原始题面“违约电价50%”'):
            p.text='原始题面“违约电价50%”未明确退款，正0.5倍下调罚费是另一种解释。本次固定退款口径，只研究终端价值与参数标定；未在新参数和新年末边界下重算另一结算解释，不能把旧口径对照混入本次结果。'
        if p.text.startswith('当前计划层的共同储能轨迹可能偏保守'):
            p.text='本版采用S=1点预测，尚未表示完整多阶段预测不确定性；线性库存价值仍是跨日近似；1月验证集不覆盖全年季节，新增候选是事后修订；整点插值依赖采样假设。上述局限不因每次优化求得最优而消失。'
    # 移除Word默认标题装饰，保持黑色学术排版。
    for st in [doc.styles['Title']]:
        pr=st.element.find(b.qn('w:pPr'))
        if pr is not None:
            border=pr.find(b.qn('w:pBdr'))
            if border is not None:pr.remove(border)
    out=m.ROOT/'第3问建模文档_终端修订版.docx';doc.save(out);print(out)
if __name__=='__main__':main()
