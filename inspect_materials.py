from pathlib import Path
import json
import openpyxl
from pypdf import PdfReader
from docx import Document

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'C题分析交付'/'材料核验'
OUT.mkdir(parents=True,exist_ok=True)
HOME=Path('C:/Users/ahdezneytiry-boop1')
for p in (HOME/'Desktop/附件').rglob('*.xlsx'):
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    print('\nFILE',p.name)
    for s in w:
        print('SHEET',s.title,s.max_row,s.max_column)
        for row in list(s.values)[:5]: print(row[:12])
refs=Path('D:/weixing/xwechat_files/wxid_jtqt6ar7kbkt12_b5f7/msg/file/2026-09')
for p in [HOME/'Desktop/C题.pdf', refs/'赛题分析(1).pdf',refs/'论文(1).pdf']:
    r=PdfReader(p)
    txt='\n'.join(f'\n第{i+1}页\n'+(page.extract_text() or '') for i,page in enumerate(r.pages))
    (OUT/(p.stem+'.txt')).write_text(txt,encoding='utf-8')
    print('PDF',p.name,len(r.pages),len(txt))
d=Document(refs/'思路-2026C(1).docx')
txt='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
(OUT/'思路.txt').write_text(txt,encoding='utf-8')
try:
    import fitz
    d=fitz.open(HOME/'Desktop/C题.pdf')
    for i,p in enumerate(d):p.get_pixmap(matrix=fitz.Matrix(1.3,1.3)).save(str(OUT/f'题目_{i+1}.png'))
    print('RENDER OK')
except ImportError:print('NO FITZ')
